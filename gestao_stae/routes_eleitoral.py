from flask import Blueprint, render_template, request, session, redirect, url_for, flash, jsonify, send_file
import sqlite3
import datetime
import os
import uuid as _uuid
try:
    import psycopg2
    import psycopg2.extras
except ImportError:
    pass

# Define o Blueprint
eleitoral_bp = Blueprint('eleitoral', __name__, url_prefix='/eleitoral')

# Estados do fluxo de distribuição (ordem de avanço)
ESTADOS_DISTRIBUICAO_ORDEM = {
    'EM_PREPARACAO': 1,
    'EMPACOTAMENTO': 2,
    'A_ESPERA_ENVIO': 3,
    'ENVIADO': 4,
    'RECEBIDO': 5,
    # Estado legado (guias antigas) — equiparado a "em trânsito" (entre envio e receção)
    'EM_TRANSITO': 4,
}
ESTADOS_DISTRIBUICAO = list(ESTADOS_DISTRIBUICAO_ORDEM.keys())[:5]

# Tipos possíveis para os locais de armazenamento (hierarquia eleitoral, sem brigadas)
ELEITORAL_TIPOS_LOCAL = [
    'DIRECAO_GERAL',
    'DIRECAO_NACIONAL',
    'DIRECAO_PROVINCIAL',
    'DEPARTAMENTO_NACIONAL',
    'DEPARTAMENTO_PROVINCIAL',
    'REPARTICAO',
    'DIRECAO_DISTRITAL',
    'POSTO_RECENSEAMENTO',
    'POSTO_VOTACAO',
    'ENTIDADE_EXTERNA',
    'CENTRAL',
    'PROVINCIA',
    'PAIS_DIASPORA',
]

TITULO_ESTADO = {
    'EM_PREPARACAO': 'Preparação',
    'EMPACOTAMENTO': 'Empacotamento',
    'A_ESPERA_ENVIO': 'A aguardar envio',
    'ENVIADO': 'Enviado',
    'RECEBIDO': 'Recebido',
    'EM_TRANSITO': 'Em trânsito',
    'ANULADA': 'Anulada',
}

def get_eleitoral_db():
    from app import DB_PATH, get_pg_connection, is_cloud_mode
    if is_cloud_mode():
        import psycopg2.extras
        conn = get_pg_connection()
        return conn, True # True indica que é Postgres
    else:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        return conn, False # False indica que é SQLite

def get_pg_for_dual_write():
    """Tenta obter ligação PG para dual-write. Retorna None se não houver internet."""
    try:
        from app import get_pg_connection
        conn = get_pg_connection()
        if conn:
            return conn
    except Exception:
        pass
    return None

def dual_execute(s_conn, pg_conn, sqlite_sql, pg_sql, params_sqlite, params_pg=None):
    """
    Executa uma query em SQLite E PostgreSQL simultaneamente.
    Se PG falhar, apenas regista o erro mas não bloqueia o SQLite.
    Retorna o cursor SQLite (fonte de verdade).
    """
    if params_pg is None:
        params_pg = params_sqlite
    
    # Sempre escreve no SQLite (fonte de verdade local)
    s_c = s_conn.cursor()
    s_c.execute(sqlite_sql, params_sqlite)
    
    # Tenta escrever no PG também (dual-write)
    if pg_conn:
        try:
            pg_c = pg_conn.cursor()
            pg_c.execute(pg_sql, params_pg)
            pg_conn.commit()
        except Exception as e:
            try:
                pg_conn.rollback()
            except:
                pass
            print(f"[dual-write] Aviso PG: {e}")
    
    return s_c

def check_permission():
    if 'username' not in session:
        return False
    perfil = session.get('perfil', '').lower()
    if perfil == 'admin':
        return True
    # Política por perfil: apenas técnicos com local eleitoral atribuído acedem
    # ao módulo de logística (users.eleitoral_local_id). Perfis como 'protecao'
    # não têm acesso.
    if perfil != 'tecnico':
        return False
    if session.get('eleitoral_local_id') is not None:
        return True
    return False

@eleitoral_bp.before_request
def before_request():
    if not check_permission():
        flash("Acesso negado ao Módulo Eleitoral.", "error")
        return redirect('/')

@eleitoral_bp.route('/')
def dashboard():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    # Pega processos para o filtro
    if is_pg:
        c.execute("SELECT * FROM eleitoral_processo_eleitoral ORDER BY ano DESC")
        processos = c.fetchall()
    else:
        c.execute("SELECT * FROM eleitoral_processo_eleitoral ORDER BY ano DESC")
        processos = c.fetchall()
    
    c.execute("SELECT id, nome, variante FROM eleitoral_tipo_material ORDER BY nome")
    tipos_material = c.fetchall()

    processo_id = request.args.get('processo_id')
    tipo_material_id = request.args.getlist('tipo_material_id')
    tipo_material_id = [x for x in tipo_material_id if x]
    
    conds = []
    params = []
    if processo_id:
        conds.append(f"processo_id = {'%s' if is_pg else '?'}")
        params.append(processo_id)
    if tipo_material_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(tipo_material_id))
        conds.append(f"tipo_material_id IN ({placeholders})")
        params.extend(tipo_material_id)
        
    cond_proc = "WHERE " + " AND ".join(conds) if conds else ""
    param = tuple(params)
        
    # Resumo
    c.execute(f"SELECT SUM(quantidade_total) as t, SUM(quantidade_bom) as b, SUM(quantidade_mau) as m FROM eleitoral_material_sobrante {cond_proc}", param)
    row = c.fetchone()
    if not is_pg and row:
        row = dict(row)
    resumo = {'total': row['t'] if row and row['t'] else 0, 'bom': row['b'] if row and row['b'] else 0, 'mau': row['m'] if row and row['m'] else 0}
    
    # Categorias
    # Se já filtramos por tipo_material_id, s.tipo_material_id já está no cond_proc, mas a tabela tem alias s. No cond_proc, os campos não tem alias.
    # Vamos reescrever cond_proc para aliases quando tem join
    cond_proc_join = "WHERE " + " AND ".join([c.replace("processo_id", "s.processo_id").replace("tipo_material_id", "s.tipo_material_id") for c in conds]) if conds else ""
    
    c.execute(f"""
        SELECT c.nome, SUM(s.quantidade_total) as total
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        JOIN eleitoral_categoria_material c ON t.categoria_id = c.id
        {cond_proc_join}
        GROUP BY c.nome
    """, param)
    cat_rows = c.fetchall()
    
    categorias_labels = [r['nome'] for r in cat_rows] if cat_rows else []
    categorias_data = [r['total'] for r in cat_rows] if cat_rows else []
    
    # Provincias
    cond_proc_prov = cond_proc_join
    if cond_proc_prov:
        cond_proc_prov += " AND p.id IS NOT NULL"
    else:
        cond_proc_prov = "WHERE p.id IS NOT NULL"
        
    c.execute(f"""
        SELECT p.nome as provincia, t.nome as tipo, SUM(s.quantidade_total) as total
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        LEFT JOIN eleitoral_provincia p ON p.id = l.provincia_id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        {cond_proc_prov}
        GROUP BY p.nome, t.nome
        ORDER BY p.nome
    """, param)
    prov_rows = c.fetchall()
    
    # Processar dados para o Chart.js
    provincias_list = []
    tipos_dict = {}
    
    if prov_rows:
        for r in prov_rows:
            p = r['provincia'] if not is_pg and isinstance(r, dict) else (r['provincia'] if is_pg else dict(r)['provincia'])
            if p not in provincias_list:
                provincias_list.append(p)
                
        for r in prov_rows:
            p = r['provincia'] if not is_pg and isinstance(r, dict) else (r['provincia'] if is_pg else dict(r)['provincia'])
            t = r['tipo'] if not is_pg and isinstance(r, dict) else (r['tipo'] if is_pg else dict(r)['tipo'])
            v = float(r['total'] if not is_pg and isinstance(r, dict) else (r['total'] if is_pg else dict(r)['total']))
            
            if t not in tipos_dict:
                tipos_dict[t] = {prov: 0 for prov in provincias_list}
            tipos_dict[t][p] = v
            
    # Criar paleta de cores
    cores = ['#8b5cf6', '#3b82f6', '#10b981', '#f59e0b', '#ef4444', '#ec4899', '#6366f1', '#14b8a6', '#f97316', '#84cc16']
    
    datasets_provincia = []
    c_idx = 0
    for tipo, valores in tipos_dict.items():
        data_arr = [valores[p] for p in provincias_list]
        datasets_provincia.append({
            'label': tipo,
            'data': data_arr,
            'backgroundColor': cores[c_idx % len(cores)]
        })
        c_idx += 1
    
    conn.close()
    return render_template('eleitoral/dashboard.html', processos=processos, resumo=resumo, 
                           categorias_labels=categorias_labels, categorias_data=categorias_data,
                           provincias_labels=provincias_list, datasets_provincia=datasets_provincia,
                           tipos_material=tipos_material, processo_id=processo_id, tipo_material_id=tipo_material_id)

@eleitoral_bp.route('/processos', methods=['GET', 'POST'])
def processos():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    if request.method == 'POST':
        nome = request.form.get('nome')
        tipo = request.form.get('tipo')
        ano = request.form.get('ano')
        data_inicio = request.form.get('data_inicio') or None
        data_fim = request.form.get('data_fim') or None
        
        if is_pg:
            c.execute("""
                INSERT INTO eleitoral_processo_eleitoral (nome, tipo, ano, data_inicio, data_fim)
                VALUES (%s, %s, %s, %s, %s)
            """, (nome, tipo, ano, data_inicio, data_fim))
        else:
            c.execute("""
                INSERT INTO eleitoral_processo_eleitoral (nome, tipo, ano, data_inicio, data_fim)
                VALUES (?, ?, ?, ?, ?)
            """, (nome, tipo, ano, data_inicio, data_fim))
        conn.commit()
        flash("Processo criado com sucesso!", "success")
        return redirect(url_for('eleitoral.processos'))
        
    c.execute("SELECT * FROM eleitoral_processo_eleitoral ORDER BY ano DESC")
    processos_list = c.fetchall()
    
    # Busca os eventos para cada processo
    c.execute("SELECT * FROM eleitoral_evento ORDER BY data_inicio DESC")
    eventos = c.fetchall()
    
    # Converter para listas de dicionários se for sqlite3.Row para podermos modificar
    proc_dicts = [dict(p) for p in processos_list] if not is_pg else processos_list
    ev_dicts = [dict(e) for e in eventos] if not is_pg else eventos
    
    for p in proc_dicts:
        p['eventos'] = [e for e in ev_dicts if e['processo_id'] == p['id']]
        
    conn.close()
    return render_template('eleitoral/processos.html', processos=proc_dicts)

@eleitoral_bp.route('/processos/<int:id>/iniciar', methods=['POST'])
def iniciar_processo(id):
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    c.execute(f"UPDATE eleitoral_processo_eleitoral SET estado = 'EM_CURSO' WHERE id = {'%s' if is_pg else '?'}", (id,))
    conn.commit()
    conn.close()
    flash("Processo iniciado!", "success")
    return redirect(url_for('eleitoral.processos'))

@eleitoral_bp.route('/processos/<int:id>/encerrar', methods=['POST'])
def encerrar_processo(id):
    destino_id = request.form.get('destino_id')
    if not destino_id:
        flash("ID do processo destino é obrigatório.", "error")
        return redirect(url_for('eleitoral.processos'))
    
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    try:
        utilizador_id = session.get('user_id') or 1
        
        # Inserir transferencia
        if is_pg:
            c.execute("INSERT INTO eleitoral_transferencia_sobrantes (processo_origem_id, processo_destino_id, utilizador_id) VALUES (%s, %s, %s) RETURNING id", (id, destino_id, utilizador_id))
            transf_id = c.fetchone()[0]
        else:
            c.execute("INSERT INTO eleitoral_transferencia_sobrantes (processo_origem_id, processo_destino_id, utilizador_id) VALUES (?, ?, ?)", (id, destino_id, utilizador_id))
            transf_id = c.lastrowid
            
        # Copiar para o historico
        param_marker = "%s" if is_pg else "?"
        c.execute(f"INSERT INTO eleitoral_transferencia_sobrantes_item (transferencia_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau) SELECT {param_marker}, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau FROM eleitoral_material_sobrante WHERE processo_id = {param_marker}", (transf_id, id))
        
        # Gerar o stock do proximo processo
        c.execute(f"INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, processo_origem_id, utilizador_id) SELECT {param_marker}, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, 'TRANSFERENCIA_SOBRANTE', {param_marker}, {param_marker} FROM eleitoral_material_sobrante WHERE processo_id = {param_marker}", (destino_id, id, utilizador_id, id))
        
        # Atualizar estados
        c.execute(f"UPDATE eleitoral_processo_eleitoral SET estado = 'ENCERRADO' WHERE id = {param_marker}", (id,))
        c.execute(f"UPDATE eleitoral_processo_eleitoral SET estado = 'EM_CURSO' WHERE id = {param_marker}", (destino_id,))
        
        conn.commit()
        flash("Processo encerrado e sobrantes transferidos!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao transferir: {str(e)}", "error")
    finally:
        conn.close()
        
    return redirect(url_for('eleitoral.processos'))

@eleitoral_bp.route('/material')
def material():
    conn, is_pg = get_eleitoral_db()
    import psycopg2.extras
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    # Processos para a dropdown
    c.execute("SELECT id, nome FROM eleitoral_processo_eleitoral ORDER BY ano DESC")
    processos = c.fetchall()
    
    # Locais para a dropdown
    c.execute("SELECT id, tipo, nome FROM eleitoral_local_armazenamento ORDER BY tipo, nome")
    locais_db = c.fetchall()
    locais = {}
    for loc in locais_db:
        locais.setdefault(loc['tipo'], []).append(loc)

    processo_id = request.args.get('processo_id')
    local_id = request.args.get('local_id')
    tipo_material_id = request.args.getlist('tipo_material_id')
    tipo_material_id = [x for x in tipo_material_id if x]
    
    materiais = []
    if processo_id:
        param_marker = "%s" if is_pg else "?"
        query_base = f"""
            SELECT m.*, t.nome as tipo_nome, t.variante as tipo_variante, l.nome as local_nome
            FROM eleitoral_material_sobrante m
            JOIN eleitoral_tipo_material t ON m.tipo_material_id = t.id
            JOIN eleitoral_local_armazenamento l ON m.local_id = l.id
            WHERE m.processo_id = {param_marker}
        """
        
        params = [processo_id]
        
        if local_id:
            query_base += f" AND m.local_id = {param_marker}"
            params.append(local_id)
            
        if tipo_material_id:
            placeholders = ','.join([param_marker] * len(tipo_material_id))
            query_base += f" AND m.tipo_material_id IN ({placeholders})"
            params.extend(tipo_material_id)
            
        c.execute(query_base, tuple(params))
        materiais = c.fetchall()
        
    c.execute("SELECT * FROM eleitoral_tipo_material ORDER BY nome")
    tipos_material = c.fetchall()
    
    conn.close()
    return render_template('eleitoral/material.html', processos=processos, locais=locais, 
                           processo_id=processo_id, local_id=local_id, tipo_material_id=tipo_material_id,
                           materiais=materiais, tipos_material=tipos_material)
@eleitoral_bp.route('/material/registar', methods=['POST'])
def registar_material():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    param_marker = "%s" if is_pg else "?"
    try:
        processo_id = request.form.get('processo_id')
        local_id = request.form.get('local_id')
        tipo_material_id = request.form.get('tipo_material_id')
        qtd_bom = float(request.form.get('quantidade_bom', 0) or 0)
        qtd_mau = float(request.form.get('quantidade_mau', 0) or 0)
        qtd_total = qtd_bom + qtd_mau
        observacoes = request.form.get('observacoes', '')
        utilizador_id = session.get('user_id', 1)

        if not processo_id or not local_id or not tipo_material_id:
            flash("Erro: preencha processo, local e material.", "error")
            return redirect(url_for('eleitoral.material'))

        # Verifica se já existe registo para processo+local+tipo — se sim, soma
        c.execute(f"SELECT id FROM eleitoral_material_sobrante WHERE processo_id = {param_marker} AND local_id = {param_marker} AND tipo_material_id = {param_marker}",
                  (processo_id, local_id, tipo_material_id))
        existing = c.fetchone()
        if existing:
            c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total + {param_marker}, quantidade_bom = quantidade_bom + {param_marker}, quantidade_mau = quantidade_mau + {param_marker}, actualizado_em = CURRENT_TIMESTAMP WHERE id = {param_marker}",
                      (qtd_total, qtd_bom, qtd_mau, existing[0]))
        else:
            c.execute(f"INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, observacoes, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, 'REGISTO_DIRECTO', {param_marker}, {param_marker})",
                      (processo_id, local_id, tipo_material_id, qtd_total, qtd_bom, qtd_mau, observacoes, utilizador_id))
        conn.commit()

        # DUAL-WRITE para PG quando em modo local
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg_c = pg.cursor()
                    pg_c.execute("SELECT id FROM eleitoral_material_sobrante WHERE processo_id=%s AND local_id=%s AND tipo_material_id=%s",
                                 (processo_id, local_id, tipo_material_id))
                    pexist = pg_c.fetchone()
                    if pexist:
                        pg_c.execute("UPDATE eleitoral_material_sobrante SET quantidade_total=quantidade_total+%s, quantidade_bom=quantidade_bom+%s, quantidade_mau=quantidade_mau+%s WHERE id=%s",
                                     (qtd_total, qtd_bom, qtd_mau, pexist[0]))
                    else:
                        pg_c.execute("INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, observacoes, utilizador_id) VALUES (%s,%s,%s,%s,%s,%s,'REGISTO_DIRECTO',%s,%s)",
                                     (processo_id, local_id, tipo_material_id, qtd_total, qtd_bom, qtd_mau, observacoes, utilizador_id))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write registar] {e}")
                    try: pg.rollback()
                    except Exception: pass
                finally:
                    pg.close()

        flash("Material registado com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao registar material: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.material', processo_id=processo_id))

@eleitoral_bp.route('/material/editar/<int:id>', methods=['POST'])
def editar_material(id):
    from flask import session
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    try:
        user_local_id = session.get('eleitoral_local_id')
        if user_local_id:
            c.execute(f"SELECT local_id FROM eleitoral_material_sobrante WHERE id = {'%s' if is_pg else '?'}", (id,))
            res = c.fetchone()
            if not res or str(res[0]) != str(user_local_id):
                from flask import flash, redirect, url_for
                flash('Acesso negado: Não tem permissões para editar este material.', 'error')
                return redirect(url_for('eleitoral.material'))

        from flask import request, flash, redirect, url_for
        qtd_bom = float(request.form.get('quantidade_bom', 0) or 0)
        qtd_mau = float(request.form.get('quantidade_mau', 0) or 0)
        qtd_total = qtd_bom + qtd_mau
        param_marker = "%s" if is_pg else "?"
        c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = {param_marker}, quantidade_bom = {param_marker}, quantidade_mau = {param_marker} WHERE id = {param_marker}", (qtd_total, qtd_bom, qtd_mau, id))
        conn.commit()
        
        # DUAL-WRITE: Se em modo local, replicar também no PG
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    # No PG, procurar pelo processo+local+tipo (IDs podem diferir)
                    c.execute(f"SELECT processo_id, local_id, tipo_material_id FROM eleitoral_material_sobrante WHERE id = ?", (id,))
                    row = c.fetchone()
                    if row:
                        pg_c = pg.cursor()
                        pg_c.execute("UPDATE eleitoral_material_sobrante SET quantidade_total=%s, quantidade_bom=%s, quantidade_mau=%s WHERE processo_id=%s AND local_id=%s AND tipo_material_id=%s",
                                     (qtd_total, qtd_bom, qtd_mau, row[0], row[1], row[2]))
                        pg.commit()
                except Exception as e:
                    print(f"[dual-write editar] {e}")
                finally:
                    pg.close()
        
        flash('Material atualizado com sucesso!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Erro: {str(e)}', 'error')
    finally:
        conn.close()
    return redirect(request.referrer or url_for('eleitoral.material'))

@eleitoral_bp.route('/material/apagar/<int:id>', methods=['POST'])
def apagar_material(id):
    from flask import session
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    try:
        user_local_id = session.get('eleitoral_local_id')
        if user_local_id:
            c.execute(f"SELECT local_id FROM eleitoral_material_sobrante WHERE id = {'%s' if is_pg else '?'}", (id,))
            res = c.fetchone()
            if not res or str(res[0]) != str(user_local_id):
                from flask import flash, redirect, url_for
                flash('Acesso negado: Não tem permissões para apagar este material.', 'error')
                return redirect(url_for('eleitoral.material'))

        from flask import request, flash, redirect, url_for
        param_marker = "%s" if is_pg else "?"
        
        # DUAL-WRITE: Guardar processo+local+tipo antes de apagar (para apagar no PG)
        pg_keys = None
        if not is_pg:
            c.execute("SELECT processo_id, local_id, tipo_material_id FROM eleitoral_material_sobrante WHERE id = ?", (id,))
            pg_keys = c.fetchone()
        
        c.execute(f"DELETE FROM eleitoral_material_sobrante WHERE id = {param_marker}", (id,))
        conn.commit()
        
        # DUAL-WRITE: Apagar também no PG
        if not is_pg and pg_keys:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg_c = pg.cursor()
                    pg_c.execute("DELETE FROM eleitoral_material_sobrante WHERE processo_id=%s AND local_id=%s AND tipo_material_id=%s",
                                 (pg_keys[0], pg_keys[1], pg_keys[2]))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write apagar] {e}")
                finally:
                    pg.close()
        
        flash('Material removido com sucesso!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Erro: {str(e)}', 'error')
    finally:
        conn.close()
    return redirect(request.referrer or url_for('eleitoral.material'))

def _registar_importacao(processo_id, caminho, nome_ficheiro, modo, utilizador_id, total_sucesso):
    """Regista uma importação em eleitoral_importacao_material (SQLite + PG via dual)."""
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    marker = "%s" if is_pg else "?"
    try:
        c.execute(f"INSERT INTO eleitoral_importacao_material (processo_id, nome_ficheiro, caminho_ficheiro, modo, utilizador_id, total_sucesso, estado) VALUES ({marker},{marker},{marker},{marker},{marker},{marker},'CONCLUIDO')",
                  (processo_id, nome_ficheiro, caminho, modo, utilizador_id, total_sucesso))
        conn.commit()
        importacao_id = c.lastrowid if not is_pg else None
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg.cursor().execute("INSERT INTO eleitoral_importacao_material (processo_id, nome_ficheiro, caminho_ficheiro, modo, utilizador_id, total_sucesso, estado) VALUES (%s,%s,%s,%s,%s,%s,'CONCLUIDO')",
                                        (processo_id, nome_ficheiro, caminho, modo, utilizador_id, total_sucesso))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write importacao-reg] {e}")
                finally:
                    pg.close()
        return importacao_id
    finally:
        conn.close()


@eleitoral_bp.route('/importacao')
def lista_importacoes():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    marker = "%s" if is_pg else "?"
    try:
        c.execute(f"SELECT i.id, i.nome_ficheiro, i.caminho_ficheiro, i.data_importacao, i.modo, i.total_sucesso, p.nome AS processo FROM eleitoral_importacao_material i LEFT JOIN eleitoral_processo_eleitoral p ON p.id = i.processo_id ORDER BY i.id DESC")
        rows = c.fetchall()
    finally:
        conn.close()
    return jsonify([dict(r) if hasattr(r, 'keys') else {
        'id': r[0], 'nome_ficheiro': r[1], 'caminho_ficheiro': r[2],
        'data_importacao': r[3], 'modo': r[4], 'total_sucesso': r[5], 'processo': r[6]
    } for r in rows])


@eleitoral_bp.route('/importacao/preview/<int:id>')
def preview_importacao(id):
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    marker = "%s" if is_pg else "?"
    c.execute(f"SELECT caminho_ficheiro FROM eleitoral_importacao_material WHERE id = {marker}", (id,))
    r = c.fetchone()
    conn.close()
    caminho = r[0] if r else None
    if not caminho or not os.path.exists(caminho):
        return jsonify({'linhas': [], 'erro': 'Ficheiro original não disponível.'})
    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho)
        sheet = wb.active
        dados = []
        for row in sheet.iter_rows(min_row=1, max_row=25, values_only=True):
            dados.append([('' if v is None else str(v)) for v in row])
        return jsonify({'linhas': dados})
    except Exception as e:
        return jsonify({'linhas': [], 'erro': f'Erro ao ler ficheiro: {str(e)}'})


@eleitoral_bp.route('/importacao', methods=['POST'])
def importacao_excel():
    if 'file' not in request.files:
        flash("Nenhum ficheiro selecionado", "error")
        return redirect(url_for('eleitoral.material'))
    
    file = request.files['file']
    if file.filename == '':
        flash("Nenhum ficheiro selecionado", "error")
        return redirect(url_for('eleitoral.material'))
        
    processo_id = request.form.get('processo_id')
    modo = request.form.get('modo_importacao', 'adicionar')
    limpar = request.form.get('limpar_provincia') == 'sim'
    
    if file and file.filename.endswith('.xlsx'):
        saved_caminho = None
        importacao_id = None
        try:
            import openpyxl

            # Guarda uma cópia do ficheiro para pré-visualização (pós-importação)
            from app import UPLOAD_FOLDER
            if UPLOAD_FOLDER:
                try:
                    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
                    st = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
                    fname = f"material_import_{st}_{os.path.basename(file.filename)}"
                    saved_caminho = os.path.join(UPLOAD_FOLDER, fname)
                    file.stream.seek(0)
                    with open(saved_caminho, 'wb') as fh:
                        fh.write(file.read())
                    file.stream.seek(0)
                except Exception as e:
                    print(f"[import] erro ao guardar ficheiro: {e}")
                    saved_caminho = None

            wb = openpyxl.load_workbook(file)
            sheet = wb.active
            
            conn, is_pg = get_eleitoral_db()
            c = conn.cursor()
            
            # Helper to fetch maps
            c.execute("SELECT id, nome FROM eleitoral_local_armazenamento")
            locais_map = {row[1].strip().lower(): row[0] for row in c.fetchall()}
            
            c.execute("SELECT id, nome, variante FROM eleitoral_tipo_material")
            tipos_map = {}
            for row in c.fetchall():
                nome = row[1].strip().lower() if row[1] else ""
                variante = row[2].strip().lower() if row[2] else ""
                tipos_map[(nome, variante)] = row[0]
                
            locais_limpos = set()
            registos_adicionados = 0
            
            # Assumimos cabeçalho na linha 1. Colunas: Provincia, Tipo, Variante, Total, Bom, Mau
            # Vamos tentar ler pelas colunas caso existam
            header = [str(cell.value).strip().lower() if cell.value else '' for cell in sheet[1]]
            
            # Fallback for indices if header doesn't match perfectly
            col_prov = header.index('província') if 'província' in header else (header.index('provincia') if 'provincia' in header else 0)
            col_tipo = header.index('tipo') if 'tipo' in header else 1
            col_var = header.index('variante') if 'variante' in header else 2
            col_tot = header.index('total') if 'total' in header else 3
            col_bom = header.index('bom') if 'bom' in header else 4
            col_mau = header.index('mau') if 'mau' in header else 5
            
            utilizador_id = session.get('user_id', 1)
            
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[col_prov] or not row[col_tipo]: continue
                
                prov = str(row[col_prov]).strip()
                tipo = str(row[col_tipo]).strip()
                var = str(row[col_var]).strip() if len(row) > col_var and row[col_var] else ""
                
                try: tot = int(row[col_tot]) if len(row) > col_tot and row[col_tot] else 0
                except: tot = 0
                try: bom = int(row[col_bom]) if len(row) > col_bom and row[col_bom] else 0
                except: bom = 0
                try: mau = int(row[col_mau]) if len(row) > col_mau and row[col_mau] else 0
                except: mau = 0
                # Garante a consistência: total = bom + mau
                tot = bom + mau
                
                # Match local
                local_id = locais_map.get(prov.lower())
                if not local_id:
                    # Tenta match parcial
                    for k, v in locais_map.items():
                        if k in prov.lower() or prov.lower() in k:
                            local_id = v
                            break
                            
                # Match tipo
                tipo_id = tipos_map.get((tipo.lower(), var.lower()))
                if not tipo_id:
                    # Tenta sem variante
                    tipo_id = tipos_map.get((tipo.lower(), ""))
                    
                if local_id and tipo_id:
                    if limpar and local_id not in locais_limpos:
                        c.execute(f"DELETE FROM eleitoral_material_sobrante WHERE processo_id = {'%s' if is_pg else '?'} AND local_id = {'%s' if is_pg else '?'}", (processo_id, local_id))
                        locais_limpos.add(local_id)
                        
                    # Verifica se já existe
                    c.execute(f"SELECT id, quantidade_total, quantidade_bom, quantidade_mau FROM eleitoral_material_sobrante WHERE processo_id = {'%s' if is_pg else '?'} AND local_id = {'%s' if is_pg else '?'} AND tipo_material_id = {'%s' if is_pg else '?'}", (processo_id, local_id, tipo_id))
                    existing = c.fetchone()
                    
                    if existing:
                        m_id = existing[0]
                        if modo == 'adicionar':
                            new_tot = existing[1] + tot
                            new_bom = existing[2] + bom
                            new_mau = existing[3] + mau
                            c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = {'%s' if is_pg else '?'}, quantidade_bom = {'%s' if is_pg else '?'}, quantidade_mau = {'%s' if is_pg else '?'} WHERE id = {'%s' if is_pg else '?'}", (new_tot, new_bom, new_mau, m_id))
                        else:
                            c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = {'%s' if is_pg else '?'}, quantidade_bom = {'%s' if is_pg else '?'}, quantidade_mau = {'%s' if is_pg else '?'} WHERE id = {'%s' if is_pg else '?'}", (tot, bom, mau, m_id))
                    else:
                        c.execute(f"INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id) VALUES ({'%s' if is_pg else '?'}, {'%s' if is_pg else '?'}, {'%s' if is_pg else '?'}, {'%s' if is_pg else '?'}, {'%s' if is_pg else '?'}, {'%s' if is_pg else '?'}, 'IMPORTACAO_EXCEL', {'%s' if is_pg else '?'})", (processo_id, local_id, tipo_id, tot, bom, mau, utilizador_id))
                    registos_adicionados += 1
            
            conn.commit()
            
            # DUAL-WRITE: Se em modo local, replicar no PG também
            if not is_pg:
                pg = get_pg_for_dual_write()
                if pg:
                    try:
                        pg_c = pg.cursor()
                        # Replicate all material for this processo
                        c.execute("SELECT processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id FROM eleitoral_material_sobrante WHERE processo_id = ?", (processo_id,))
                        for ms_row in c.fetchall():
                            pg_c.execute("""
                                INSERT INTO eleitoral_material_sobrante 
                                    (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id)
                                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                                ON CONFLICT (processo_id, local_id, tipo_material_id) 
                                DO UPDATE SET quantidade_total=EXCLUDED.quantidade_total, quantidade_bom=EXCLUDED.quantidade_bom, quantidade_mau=EXCLUDED.quantidade_mau
                            """, (ms_row[0], ms_row[1], ms_row[2], ms_row[3], ms_row[4], ms_row[5], ms_row[6], ms_row[7]))
                        pg.commit()
                        print(f"[dual-write importacao] Replicado para PG com sucesso")
                    except Exception as e:
                        print(f"[dual-write importacao] Erro PG: {e}")
                        try: pg.rollback()
                        except: pass
                    finally:
                        pg.close()
            
            conn.close()
            # Regista a importação (parcialmente para pré-visualização/fins de auditoria)
            try:
                importacao_id = _registar_importacao(
                    processo_id, saved_caminho, file.filename,
                    modo, utilizador_id, registos_adicionados
                )
            except Exception as e:
                print(f"[import] erro ao registar importacao: {e}")
            flash(f"Importação concluída! {registos_adicionados} registos processados.", "success")
        except Exception as e:
            flash(f"Erro ao processar ficheiro Excel: {str(e)}", "error")
    else:
        flash("Apenas ficheiros .xlsx são permitidos.", "error")
        
    return redirect(url_for('eleitoral.material'))

@eleitoral_bp.route('/catalogos')
def catalogos():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    c.execute("SELECT * FROM eleitoral_tipo_material")
    tipos = c.fetchall()
    c.execute("SELECT * FROM eleitoral_categoria_material")
    categorias = c.fetchall()
    c.execute("SELECT * FROM eleitoral_provincia")
    provincias = c.fetchall()
    c.execute("SELECT * FROM eleitoral_pais_diaspora")
    paises = c.fetchall()
    conn.close()
    return render_template('eleitoral/catalogos.html', tipos=tipos, categorias=categorias, provincias=provincias, paises=paises)

@eleitoral_bp.route('/categoria_material/novo', methods=['POST'])
def novo_categoria_material():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    param_marker = "%s" if is_pg else "?"
    try:
        nome = request.form.get('nome', '').strip()
        if not nome:
            flash("Erro: o nome da categoria é obrigatório.", "error")
            return redirect(url_for('eleitoral.catalogos'))
        # evita duplicados
        c.execute(f"SELECT id FROM eleitoral_categoria_material WHERE nome = {param_marker}", (nome,))
        if c.fetchone():
            flash("A categoria já existe.", "warning")
            return redirect(url_for('eleitoral.catalogos'))
        c.execute(f"INSERT INTO eleitoral_categoria_material (nome) VALUES ({param_marker})", (nome,))
        conn.commit()
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg.cursor().execute("INSERT INTO eleitoral_categoria_material (nome) VALUES (%s) ON CONFLICT (nome) DO NOTHING", (nome,))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write cat] {e}")
                finally:
                    pg.close()
        flash(f"Categoria '{nome}' criada com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.catalogos'))

@eleitoral_bp.route('/tipo_material/novo', methods=['POST'])
def novo_tipo_material():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    param_marker = "%s" if is_pg else "?"
    try:
        categoria_id = request.form.get('categoria_id') or None
        nome = request.form.get('nome', '').strip()
        variante = request.form.get('variante', '').strip()
        unidade_medida = request.form.get('unidade_medida', 'Unidade').strip()
        controla_estado = 1 if request.form.get('controla_estado') in ('1', 'on', 'sim') else 0
        if not nome:
            flash("Erro: o nome do material é obrigatório.", "error")
            return redirect(url_for('eleitoral.catalogos'))
        if is_pg:
            c.execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (%s,%s,%s,%s,%s)",
                      (categoria_id, nome, variante or None, unidade_medida, controla_estado))
        else:
            c.execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (?,?,?,?,?)",
                      (categoria_id, nome, variante or None, unidade_medida, controla_estado))
        conn.commit()
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg.cursor().execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (%s,%s,%s,%s,%s)",
                                        (categoria_id, nome, variante or None, unidade_medida, controla_estado))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write tipo] {e}")
                finally:
                    pg.close()
        flash(f"Material '{nome}' adicionado com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.catalogos'))

@eleitoral_bp.route('/tipo_material/nome', methods=['POST'])
def novo_tipo_material_texto():
    """Adiciona rapidamente um tipo de material apenas pelo nome (texto livre)."""
    nome = request.form.get('nome', '').strip()
    if not nome:
        flash("Erro: escreva o nome do material.", "error")
        return redirect(request.referrer or url_for('eleitoral.material'))
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    param_marker = "%s" if is_pg else "?"
    try:
        # Encontra/usa a categoria "Meios Circulantes" se houver, senão coloca sem categoria
        c.execute(f"SELECT id FROM eleitoral_categoria_material WHERE nome = {param_marker}", ("Meios Circulantes",))
        cat = c.fetchone()
        categoria_id = cat[0] if cat else None
        c.execute(f"SELECT id FROM eleitoral_tipo_material WHERE nome = {param_marker}", (nome,))
        if c.fetchone():
            flash("Este material já existe no catálogo.", "warning")
            return redirect(request.referrer or url_for('eleitoral.material'))
        if is_pg:
            c.execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (%s,%s,NULL,'Unidade',1)",
                      (categoria_id, nome))
        else:
            c.execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (?,?,NULL,'Unidade',1)",
                      (categoria_id, nome))
        conn.commit()
        tipo_id = c.lastrowid if not is_pg else None
        if not is_pg:
            pg = get_pg_for_dual_write()
            if pg:
                try:
                    pg.cursor().execute("INSERT INTO eleitoral_tipo_material (categoria_id, nome, variante, unidade_medida, controla_estado) VALUES (%s,%s,NULL,'Unidade',1)",
                                        (categoria_id, nome))
                    pg.commit()
                except Exception as e:
                    print(f"[dual-write tipo texto] {e}")
                finally:
                    pg.close()
        flash(f"Material '{nome}' adicionado ao catálogo!", "success")
        return redirect(url_for('eleitoral.material', tipo_material_id=tipo_id) if tipo_id else url_for('eleitoral.material'))
    except Exception as e:
        conn.rollback()
        flash(f"Erro: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(request.referrer or url_for('eleitoral.material'))


@eleitoral_bp.route('/eventos', methods=['POST'])
def add_evento():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    processo_id = request.form.get('processo_id')
    tipo = request.form.get('tipo')
    nome = request.form.get('nome')
    data_inicio = request.form.get('data_inicio') or None
    data_fim = request.form.get('data_fim') or None

    if is_pg:
        c.execute("""
            INSERT INTO eleitoral_evento (processo_id, tipo, nome, data_inicio, data_fim)
            VALUES (%s, %s, %s, %s, %s)
        """, (processo_id, tipo, nome, data_inicio, data_fim))
    else:
        c.execute("""
            INSERT INTO eleitoral_evento (processo_id, tipo, nome, data_inicio, data_fim)
            VALUES (?, ?, ?, ?, ?)
        """, (processo_id, tipo, nome, data_inicio, data_fim))
    
    conn.commit()
    conn.close()
    flash(f"Evento {nome} criado com sucesso!", "success")
    return redirect(url_for('eleitoral.processos'))

@eleitoral_bp.route('/eventos/<int:id>/estado', methods=['POST'])
def update_estado_evento(id):
    estado = request.form.get('estado')
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    # Se formos activar um evento, podemos querer passar os outros do mesmo processo a CONCLUIDO
    # Mas por agora apenas actualiza o estado
    param = (estado, id) if is_pg else (estado, id)
    marker = "%s" if is_pg else "?"
    c.execute(f"UPDATE eleitoral_evento SET estado = {marker} WHERE id = {marker}", param)
    conn.commit()
    conn.close()
    flash(f"Estado do evento actualizado para {estado}!", "success")
    return redirect(url_for('eleitoral.processos'))

import io
from flask import send_file

@eleitoral_bp.route('/relatorios/exportar_excel')
def exportar_excel():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    processo_id = request.args.get('processo_id')
    tipo_material_id = request.args.getlist('tipo_material_id')
    tipo_material_id = [x for x in tipo_material_id if x]
    
    conds = []
    params = []
    if processo_id:
        conds.append(f"s.processo_id = {'%s' if is_pg else '?'}")
        params.append(processo_id)
    if tipo_material_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(tipo_material_id))
        conds.append(f"s.tipo_material_id IN ({placeholders})")
        params.extend(tipo_material_id)
        
    cond_proc = "WHERE " + " AND ".join(conds) if conds else ""
    param = tuple(params)
    
    # Por Categoria
    c.execute(f'''
        SELECT c.nome as categoria, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        JOIN eleitoral_categoria_material c ON t.categoria_id = c.id
        {cond_proc}
        GROUP BY c.nome, t.nome, t.variante
        ORDER BY c.nome, total DESC
    ''', param)
    relatorio_categoria = c.fetchall()
    
    # Por Local
    c.execute(f'''
        SELECT l.nome as provincia, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        {cond_proc}
        GROUP BY l.nome, t.nome, t.variante
        ORDER BY l.nome, total DESC
    ''', param)
    relatorio_provincia = c.fetchall()
    
    conn.close()
    
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    
    wb = Workbook()
    
    # Estilos
    header_fill = PatternFill(start_color="FF1E293B", end_color="FF1E293B", fill_type="solid")
    header_font = Font(color="FFFFFFFF", bold=True)
    
    bom_fill = PatternFill(start_color="FFDCFCE7", end_color="FFDCFCE7", fill_type="solid")
    bom_font = Font(color="FF166534", bold=True)
    
    mau_fill = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
    mau_font = Font(color="FF991B1B", bold=True)
    
    # Aba 1: Categoria
    ws1 = wb.active
    ws1.title = "Por Categoria"
    headers = ["Categoria", "Tipo de Material", "Total", "Bom", "Mau"]
    ws1.append(headers)
    for col in range(1, 6):
        cell = ws1.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        
    for r in relatorio_categoria:
        tipo_str = f"{r['tipo' if is_pg else 'tipo']}"
        if r['variante' if is_pg else 'variante']:
            tipo_str += f" ({r['variante' if is_pg else 'variante']})"
            
        ws1.append([
            r['categoria' if is_pg else 'categoria'],
            tipo_str,
            r['total' if is_pg else 'total'],
            r['bom' if is_pg else 'bom'],
            r['mau' if is_pg else 'mau']
        ])
        
        # Colorir células Bom e Mau
        ws1.cell(row=ws1.max_row, column=4).fill = bom_fill
        ws1.cell(row=ws1.max_row, column=4).font = bom_font
        ws1.cell(row=ws1.max_row, column=5).fill = mau_fill
        ws1.cell(row=ws1.max_row, column=5).font = mau_font

    # Aba 2: Província
    ws2 = wb.create_sheet(title="Por Local")
    ws2.append(["Local", "Tipo de Material", "Total", "Bom", "Mau"])
    for col in range(1, 6):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        
    for r in relatorio_provincia:
        tipo_str = f"{r['tipo' if is_pg else 'tipo']}"
        if r['variante' if is_pg else 'variante']:
            tipo_str += f" ({r['variante' if is_pg else 'variante']})"
            
        ws2.append([
            r['provincia' if is_pg else 'provincia'],
            tipo_str,
            r['total' if is_pg else 'total'],
            r['bom' if is_pg else 'bom'],
            r['mau' if is_pg else 'mau']
        ])
        
        # Colorir células Bom e Mau
        ws2.cell(row=ws2.max_row, column=4).fill = bom_fill
        ws2.cell(row=ws2.max_row, column=4).font = bom_font
        ws2.cell(row=ws2.max_row, column=5).fill = mau_fill
        ws2.cell(row=ws2.max_row, column=5).font = mau_font
        
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(output, download_name="relatorio_estatistico.xlsx", as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@eleitoral_bp.route('/relatorios/exportar_pdf')
def exportar_pdf():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()

    processo_id = request.args.get('processo_id')
    tipo_material_id = request.args.getlist('tipo_material_id')
    tipo_material_id = [x for x in tipo_material_id if x]
    categoria_id = request.args.getlist('categoria_id')
    categoria_id = [x for x in categoria_id if x]
    local_id = request.args.getlist('local_id')
    local_id = [x for x in local_id if x]

    conds = []
    params = []
    if processo_id:
        conds.append(f"s.processo_id = {'%s' if is_pg else '?'}")
        params.append(processo_id)
    if tipo_material_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(tipo_material_id))
        conds.append(f"s.tipo_material_id IN ({placeholders})")
        params.extend(tipo_material_id)
    if categoria_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(categoria_id))
        conds.append(f"t.categoria_id IN ({placeholders})")
        params.extend(categoria_id)
    if local_id and not (len(local_id) == 1 and local_id[0] == ''):
        placeholders = ','.join(['%s' if is_pg else '?'] * len(local_id))
        conds.append(f"s.local_id IN ({placeholders})")
        params.extend(local_id)

    cond_proc = "WHERE " + " AND ".join(conds) if conds else ""
    param = tuple(params)

    desc_processo = "Todos os Processos"
    if processo_id:
        c.execute(f"SELECT nome, ano FROM eleitoral_processo_eleitoral WHERE id = {'%s' if is_pg else '?'}", (processo_id,))
        pr = c.fetchone()
        if pr:
            desc_processo = f"{pr['nome' if is_pg else 'nome']} ({pr['ano' if is_pg else 'ano']})"

    # Totais globais
    c.execute(f"SELECT COALESCE(SUM(quantidade_total),0) as t, COALESCE(SUM(quantidade_bom),0) as b, COALESCE(SUM(quantidade_mau),0) as m FROM eleitoral_material_sobrante s {cond_proc}", param)
    row = c.fetchone()
    totais = {
        'total': row['t' if is_pg else 't'] or 0,
        'bom': row['b' if is_pg else 'b'] or 0,
        'mau': row['m' if is_pg else 'm'] or 0
    }

    # Por categoria
    c.execute(f'''
        SELECT c.nome as categoria, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        JOIN eleitoral_categoria_material c ON t.categoria_id = c.id
        {cond_proc}
        GROUP BY c.nome, t.nome, t.variante
        ORDER BY c.nome, total DESC
    ''', param)
    relatorio_categoria = c.fetchall()

    # Por local
    c.execute(f'''
        SELECT l.nome as provincia, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        {cond_proc}
        GROUP BY l.nome, t.nome, t.variante
        ORDER BY l.nome, total DESC
    ''', param)
    relatorio_provincia = c.fetchall()

    conn.close()

    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('Titulo', parent=styles['Title'], fontSize=16, alignment=1, spaceAfter=6)
    sub_style = ParagraphStyle('Sub', parent=styles['Normal'], fontSize=11, alignment=1, spaceAfter=12, textColor=colors.HexColor('#475569'))
    head_style = ParagraphStyle('Hd', parent=styles['Normal'], fontSize=9, fontWeight='bold', textColor=colors.white)

    elementos = []

    elementos.append(Paragraph("REPÚBLICA DE MOÇAMBIQUE", title_style))
    elementos.append(Paragraph("STAE — Gestão Eleitoral · Relatório Estatístico de Material Sobrante", sub_style))
    elementos.append(Paragraph(f"<b>Processo:</b> {desc_processo}", styles['Normal']))
    if tipo_material_id:
        elementos.append(Paragraph(f"<b>Tipos de Material selecionados:</b> {len(tipo_material_id)}", styles['Normal']))
    if categoria_id:
        elementos.append(Paragraph(f"<b>Categorias selecionadas:</b> {len(categoria_id)}", styles['Normal']))
    if local_id and not (len(local_id) == 1 and local_id[0] == ''):
        elementos.append(Paragraph(f"<b>Locais selecionados:</b> {len(local_id)}", styles['Normal']))
    elementos.append(Spacer(1, 8))

    # Resumo
    resumo_data = [
        [Paragraph("Total Global", head_style), Paragraph("Bom Estado", head_style), Paragraph("Mau Estado", head_style)],
        [str(totais['total']), str(totais['bom']), str(totais['mau'])]
    ]
    resumo_tab = Table(resumo_data, colWidths=[doc.width/3.0]*3)
    resumo_tab.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.black),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('FONTSIZE', (0,0), (-1,-1), 11),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#f1f5f9')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    elementos.append(resumo_tab)
    elementos.append(Spacer(1, 14))

    # Tabela por categoria
    if relatorio_categoria:
        elementos.append(Paragraph("Agregação por Categoria", styles['Heading2']))
        data = [[Paragraph("Categoria", head_style), Paragraph("Tipo de Material", head_style),
                 Paragraph("Total", head_style), Paragraph("Bom", head_style), Paragraph("Mau", head_style)]]
        for r in relatorio_categoria:
            tipo_str = f"{r['tipo' if is_pg else 'tipo']}"
            if r['variante' if is_pg else 'variante']:
                tipo_str += f" ({r['variante' if is_pg else 'variante']})"
            data.append([
                f"{r['categoria' if is_pg else 'categoria']}",
                tipo_str,
                f"{r['total' if is_pg else 'total']}",
                f"{r['bom' if is_pg else 'bom']}",
                f"{r['mau' if is_pg else 'mau']}"
            ])
        tab = Table(data, repeatRows=1)
        tab.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
            ('ALIGN', (2,0), (-1,-1), 'RIGHT'),
        ]))
        elementos.append(tab)
        elementos.append(Spacer(1, 12))

    # Tabela por local
    if relatorio_provincia:
        elementos.append(Paragraph("Totais por Local de Armazenamento", styles['Heading2']))
        data = [[Paragraph("Local", head_style), Paragraph("Tipo de Material", head_style),
                 Paragraph("Total", head_style), Paragraph("Bom", head_style), Paragraph("Mau", head_style)]]
        for r in relatorio_provincia:
            tipo_str = f"{r['tipo' if is_pg else 'tipo']}"
            if r['variante' if is_pg else 'variante']:
                tipo_str += f" ({r['variante' if is_pg else 'variante']})"
            data.append([
                f"{r['provincia' if is_pg else 'provincia']}",
                tipo_str,
                f"{r['total' if is_pg else 'total']}",
                f"{r['bom' if is_pg else 'bom']}",
                f"{r['mau' if is_pg else 'mau']}"
            ])
        tab = Table(data, repeatRows=1)
        tab.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
            ('ALIGN', (2,0), (-1,-1), 'RIGHT'),
        ]))
        elementos.append(tab)

    doc.build(elementos)
    output.seek(0)
    return send_file(output, download_name="relatorio_estatistico.pdf", as_attachment=True, mimetype="application/pdf")


@eleitoral_bp.route('/relatorios')
def relatorios():
    conn, is_pg = get_eleitoral_db()
    import psycopg2.extras
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    c.execute("SELECT * FROM eleitoral_processo_eleitoral ORDER BY ano DESC")
    processos = c.fetchall()
    
    c.execute("SELECT id, nome, variante FROM eleitoral_tipo_material ORDER BY nome")
    tipos_material = c.fetchall()

    c.execute("SELECT id, nome FROM eleitoral_categoria_material ORDER BY nome")
    categorias = c.fetchall()

    c.execute("SELECT id, nome FROM eleitoral_local_armazenamento ORDER BY nome")
    locais = c.fetchall()

    processo_id = request.args.get('processo_id')
    tipo_material_id = request.args.getlist('tipo_material_id')
    tipo_material_id = [x for x in tipo_material_id if x]
    categoria_id = request.args.getlist('categoria_id')
    categoria_id = [x for x in categoria_id if x]
    local_id = request.args.getlist('local_id')
    local_id = [x for x in local_id if x]

    conds = []
    params = []
    if processo_id:
        conds.append(f"s.processo_id = {'%s' if is_pg else '?'}")
        params.append(processo_id)
    if tipo_material_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(tipo_material_id))
        conds.append(f"s.tipo_material_id IN ({placeholders})")
        params.extend(tipo_material_id)
    if categoria_id:
        placeholders = ','.join(['%s' if is_pg else '?'] * len(categoria_id))
        conds.append(f"t.categoria_id IN ({placeholders})")
        params.extend(categoria_id)
    if local_id and not (len(local_id) == 1 and local_id[0] == ''):
        placeholders = ','.join(['%s' if is_pg else '?'] * len(local_id))
        conds.append(f"s.local_id IN ({placeholders})")
        params.extend(local_id)
        
    cond_proc = "WHERE " + " AND ".join(conds) if conds else ""
    param = tuple(params)
        
    # Totais Globais
    c.execute(f"SELECT SUM(quantidade_total) as t, SUM(quantidade_bom) as b, SUM(quantidade_mau) as m FROM eleitoral_material_sobrante s JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id {cond_proc}", param)
    row = c.fetchone()
    if not is_pg and row:
        row = dict(row)
    resumo = {'total': row['t'] if row and row['t'] else 0, 'bom': row['b'] if row and row['b'] else 0, 'mau': row['m'] if row and row['m'] else 0}
    
    # Por Categoria
    c.execute(f"""
        SELECT c.nome as categoria, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        JOIN eleitoral_categoria_material c ON t.categoria_id = c.id
        {cond_proc}
        GROUP BY c.nome, t.nome, t.variante
        ORDER BY c.nome, total DESC
    """, param)
    relatorio_categoria = c.fetchall()
    
    # Por Tipo de Material
    c.execute(f"""
        SELECT t.nome as tipo, t.variante as variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        {cond_proc}
        GROUP BY t.nome, t.variante
        ORDER BY total DESC
    """, param)
    relatorio_tipo = c.fetchall()

    # Por Local de Armazenamento (Geográfico/Central)
    cond_proc_prov = cond_proc
        
    c.execute(f"""
        SELECT l.nome as provincia, t.nome as tipo, t.variante, SUM(s.quantidade_total) as total, SUM(s.quantidade_bom) as bom, SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        {cond_proc_prov}
        GROUP BY l.nome, t.nome, t.variante
        ORDER BY l.nome, total DESC
    """, param)
    relatorio_provincia = c.fetchall()
    
    conn.close()
    
    return render_template('eleitoral/relatorios.html', processos=processos, 
                           processo_id=processo_id, tipo_material_id=tipo_material_id, 
                           tipos_material=tipos_material,
                           categoria_id=categoria_id, categorias=categorias,
                           local_id=local_id, locais=locais,
                           resumo=resumo, relatorio_categoria=relatorio_categoria, 
                           relatorio_tipo=relatorio_tipo, relatorio_provincia=relatorio_provincia)

@eleitoral_bp.route('/distribuicao/<int:id>/confirmar', methods=['POST'])
def confirmar_rececao(id):
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    utilizador_id = session.get('user_id')
    param_marker = '%s' if is_pg else '?'
    
    try:
        # Get movimento details
        c.execute(f"SELECT * FROM eleitoral_movimento_material WHERE id = {param_marker}", (id,))
        mov = c.fetchone()
        
        if not mov or mov['estado' if is_pg else 'estado'] not in ('ENVIADO', 'EM_TRANSITO'):
            conn.close()
            flash("Movimento inválido ou ainda não enviado.", "error")
            return redirect(url_for('eleitoral.distribuicao'))
            
        processo_id = mov['processo_id' if is_pg else 'processo_id']
        destino_id = mov['local_destino_id' if is_pg else 'local_destino_id']
        
        user_perfil = session.get('perfil')
        user_local_id = session.get('eleitoral_local_id')
        
        if user_perfil != 'admin' and user_local_id:
            c.execute(f"SELECT parent_id FROM eleitoral_local_armazenamento WHERE id = {param_marker}", (destino_id,))
            dest_local = c.fetchone()
            if not dest_local or (destino_id != user_local_id and dest_local['parent_id' if is_pg else 'parent_id'] != user_local_id):
                conn.close()
                flash("Sem permissões para rececionar material neste destino.", "error")
                return redirect(url_for('eleitoral.distribuicao'))
                
        # Update estado
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        c.execute(f"UPDATE eleitoral_movimento_material SET estado = 'RECEBIDO', data_recepcao = {param_marker}, utilizador_recepcao_id = {param_marker} WHERE id = {param_marker}", (now, utilizador_id, id))
        # Regista a receção no histórico
        c.execute(f"INSERT INTO eleitoral_movimento_historico (movimento_id, estado, observacoes, data, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})",
                  (id, 'RECEBIDO', 'Receção do material confirmada pelo destino', now, utilizador_id))
        
        # Get itens
        c.execute(f"SELECT * FROM eleitoral_movimento_item WHERE movimento_id = {param_marker}", (id,))
        itens = c.fetchall()
        
        for item in itens:
            tm_id = item['tipo_material_id' if is_pg else 'tipo_material_id']
            q_bom = item['quantidade_bom' if is_pg else 'quantidade_bom']
            q_mau = item['quantidade_mau' if is_pg else 'quantidade_mau']
            q_total = q_bom + q_mau
            
            # Check if exists in Destino
            c.execute(f"SELECT id FROM eleitoral_material_sobrante WHERE processo_id = {param_marker} AND local_id = {param_marker} AND tipo_material_id = {param_marker}", (processo_id, destino_id, tm_id))
            row = c.fetchone()
            
            if row:
                row_id = row['id' if is_pg else 'id']
                c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total + {param_marker}, quantidade_bom = quantidade_bom + {param_marker}, quantidade_mau = quantidade_mau + {param_marker} WHERE id = {param_marker}",
                          (q_total, q_bom, q_mau, row_id))
            else:
                c.execute(f"INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, 'MOVIMENTO', {param_marker})",
                          (processo_id, destino_id, tm_id, q_total, q_bom, q_mau, utilizador_id))
                          
        conn.commit()
        flash("Guia rececionada com sucesso! Material somado ao stock de destino.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao confirmar receção: {str(e)}", "error")
        
    conn.close()
    return redirect(url_for('eleitoral.distribuicao'))


@eleitoral_bp.route('/distribuicao/<int:id>/anular', methods=['POST'])
def anular_movimento(id):
    if session.get('perfil') != 'admin':
        flash("Apenas administradores podem anular movimentações.", "error")
        return redirect(url_for('eleitoral.distribuicao'))

    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    param_marker = '%s' if is_pg else '?'
    utilizador_id = session.get('user_id')

    try:
        c.execute(f"SELECT * FROM eleitoral_movimento_material WHERE id = {param_marker}", (id,))
        mov = c.fetchone()

        if not mov:
            conn.close()
            flash("Guia não encontrada.", "error")
            return redirect(url_for('eleitoral.distribuicao'))

        estado = mov['estado' if is_pg else 'estado']
        if estado == 'ANULADA':
            conn.close()
            flash("Guia já se encontra anulada.", "error")
            return redirect(url_for('eleitoral.distribuicao'))

        processo_id = mov['processo_id' if is_pg else 'processo_id']
        origem_id = mov['local_origem_id' if is_pg else 'local_origem_id']
        destino_id = mov['local_destino_id' if is_pg else 'local_destino_id']
        foi_recebida = estado == 'RECEBIDO'

        c.execute(f"SELECT * FROM eleitoral_movimento_item WHERE movimento_id = {param_marker}", (id,))
        itens = c.fetchall()

        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        observacoes = (request.form.get('observacoes') or '').strip() or 'Guia anulada pelo administrador'

        for item in itens:
            tm_id = item['tipo_material_id' if is_pg else 'tipo_material_id']
            q_bom = item['quantidade_bom' if is_pg else 'quantidade_bom']
            q_mau = item['quantidade_mau' if is_pg else 'quantidade_mau']
            q_total = q_bom + q_mau

            # 1. Repor quantidades na origem (inverte a dedução feita ao emitir a guia)
            c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total + {param_marker}, quantidade_bom = quantidade_bom + {param_marker}, quantidade_mau = quantidade_mau + {param_marker} WHERE processo_id = {param_marker} AND local_id = {param_marker} AND tipo_material_id = {param_marker}",
                      (q_total, q_bom, q_mau, processo_id, origem_id, tm_id))

            # 2. Se a guia já foi recebida, retira as quantidades do destino (inverte a soma da receção)
            if foi_recebida:
                c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total - {param_marker}, quantidade_bom = quantidade_bom - {param_marker}, quantidade_mau = quantidade_mau - {param_marker} WHERE processo_id = {param_marker} AND local_id = {param_marker} AND tipo_material_id = {param_marker}",
                          (q_total, q_bom, q_mau, processo_id, destino_id, tm_id))

        # 3. Marcar a guia como anulada e registar no histórico
        c.execute(f"UPDATE eleitoral_movimento_material SET estado = 'ANULADA' WHERE id = {param_marker}", (id,))
        c.execute(f"INSERT INTO eleitoral_movimento_historico (movimento_id, estado, observacoes, data, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})",
                  (id, 'ANULADA', observacoes, now, utilizador_id))

        conn.commit()
        if foi_recebida:
            flash("Guia anulada! Quantidades repostas na origem e retiradas do destino.", "success")
        else:
            flash("Guia anulada! Quantidades repostas na origem.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao anular guia: {str(e)}", "error")

    conn.close()
    return redirect(url_for('eleitoral.distribuicao'))


@eleitoral_bp.route('/api/mapa_distribuicao')
def api_mapa_distribuicao():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    processo_id = request.args.get('processo_id')
    # Sem processo selecionado => agrega TODOS os processos (consistente com os gráficos)
    filtro_proc = ""
    param = ()
    if processo_id:
        try:
            processo_id = int(processo_id)
        except (TypeError, ValueError):
            processo_id = None
        if processo_id:
            filtro_proc = "s.processo_id = %s AND " if is_pg else "s.processo_id = ? AND "
            param = (processo_id,)

    # Map by provincia: agrega pelo nome do LOCAL (que é o nome da província),
    # excluindo Centrais e Países de Diáspora (não são províncias do mapa)
    query = '''
        SELECT 
            l.nome as provincia, 
            SUM(s.quantidade_total) as total,
            SUM(s.quantidade_bom) as bom,
            SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        WHERE {filtro}l.tipo = 'PROVINCIA'
        GROUP BY l.nome
    '''.format(filtro=filtro_proc) if is_pg else '''
        SELECT 
            l.nome as provincia, 
            SUM(s.quantidade_total) as total,
            SUM(s.quantidade_bom) as bom,
            SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        WHERE {filtro}l.tipo = 'PROVINCIA'
        GROUP BY l.nome
    '''.format(filtro=filtro_proc)
    c.execute(query, param)
    dados = c.fetchall()
    
    conn.close()
    
    resultado = {}
    for d in dados:
        if d['provincia' if is_pg else 'provincia']:
            resultado[d['provincia' if is_pg else 'provincia']] = {
                'total': d['total' if is_pg else 'total'],
                'bom': d['bom' if is_pg else 'bom'],
                'mau': d['mau' if is_pg else 'mau']
            }
            
    return jsonify(resultado)


@eleitoral_bp.route('/api/mapa_provincia')
def api_mapa_provincia():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    pm = '%s' if is_pg else '?'

    provincia = (request.args.get('provincia') or '').strip()
    processo_id = request.args.get('processo_id') or None

    if not provincia:
        conn.close()
        return jsonify({'error': 'provincia em falta', 'itens': []})

    if processo_id:
        try:
            processo_id = int(processo_id)
        except (TypeError, ValueError):
            processo_id = None
    # Sem processo selecionado => agrega TODOS os processos (consistente com os gráficos)
    filtro_proc = ""
    params = (provincia,)
    if processo_id:
        filtro_proc = f"s.processo_id = {pm} AND "
        params = (processo_id, provincia)

    query = f'''
        SELECT l.nome as provincia,
               SUM(s.quantidade_total) as total,
               SUM(s.quantidade_bom) as bom,
               SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        WHERE {filtro_proc}l.tipo = 'PROVINCIA' AND l.nome = {pm}
        GROUP BY l.nome
    '''
    c.execute(query, params)
    resumo = c.fetchone()

    itens_query = f'''
        SELECT t.nome as tipo, t.variante as variante, c.nome as categoria,
               l.nome as local, 
               SUM(s.quantidade_total) as total,
               SUM(s.quantidade_bom) as bom,
               SUM(s.quantidade_mau) as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        LEFT JOIN eleitoral_categoria_material c ON t.categoria_id = c.id
        WHERE {filtro_proc}l.tipo = 'PROVINCIA' AND l.nome = {pm}
        GROUP BY t.nome, t.variante, c.nome, l.nome
        ORDER BY total DESC
    '''
    c.execute(itens_query, params)
    itens = c.fetchall()

    conn.close()

    def g(d, k):
        return d[k if is_pg else k]

    return jsonify({
        'provincia': provincia,
        'total': g(resumo, 'total') if resumo else 0,
        'bom': g(resumo, 'bom') if resumo else 0,
        'mau': g(resumo, 'mau') if resumo else 0,
        'itens': [{
            'tipo': g(it, 'tipo'),
            'variante': g(it, 'variante'),
            'categoria': g(it, 'categoria'),
            'local': g(it, 'local'),
            'total': g(it, 'total'),
            'bom': g(it, 'bom'),
            'mau': g(it, 'mau')
        } for it in itens]
    })


@eleitoral_bp.route('/distribuicao/<int:id>/estado', methods=['POST'])
def atualizar_estado_movimento(id):
    novo_estado = request.form.get('estado')
    observacoes = (request.form.get('observacoes') or '').strip()
    
    if novo_estado not in ESTADOS_DISTRIBUICAO:
        flash("Estado inválido.", "error")
        return redirect(url_for('eleitoral.distribuicao'))
    if novo_estado == 'RECEBIDO':
        flash("Para confirmar a receção use o botão 'Receber' na guia (soma o stock ao destino).", "error")
        return redirect(url_for('eleitoral.distribuicao'))
    
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    param_marker = '%s' if is_pg else '?'
    utilizador_id = session.get('user_id')
    
    try:
        c.execute(f"SELECT * FROM eleitoral_movimento_material WHERE id = {param_marker}", (id,))
        mov = c.fetchone()
        if not mov:
            conn.close()
            flash("Guia não encontrada.", "error")
            return redirect(url_for('eleitoral.distribuicao'))
        
        estado_atual = mov['estado' if is_pg else 'estado']
        if estado_atual == 'ANULADA':
            conn.close()
            flash("Guia anulada não pode avançar de estado.", "error")
            return redirect(url_for('eleitoral.distribuicao'))
        idx_novo = ESTADOS_DISTRIBUICAO_ORDEM[novo_estado]
        idx_atual = ESTADOS_DISTRIBUICAO_ORDEM.get(estado_atual, 0)
        
        if idx_novo <= idx_atual:
            conn.close()
            flash("Só é possível avançar para um estado superior ao atual.", "error")
            return redirect(url_for('eleitoral.distribuicao'))
        
        user_perfil = session.get('perfil')
        user_local_id = session.get('eleitoral_local_id')
        origem_id = mov['local_origem_id' if is_pg else 'local_origem_id']
        
        if user_perfil != 'admin':
            if not user_local_id or origem_id != user_local_id:
                conn.close()
                flash("Apenas o local de origem pode avançar o estado da guia.", "error")
                return redirect(url_for('eleitoral.distribuicao'))
        
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        c.execute(f"UPDATE eleitoral_movimento_material SET estado = {param_marker} WHERE id = {param_marker}", (novo_estado, id))
        c.execute(f"INSERT INTO eleitoral_movimento_historico (movimento_id, estado, observacoes, data, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})",
                  (id, novo_estado, observacoes or f"Avanço para {TITULO_ESTADO.get(novo_estado, novo_estado)}", now, utilizador_id))
        
        conn.commit()
        flash(f"Guia avançada para: {TITULO_ESTADO.get(novo_estado, novo_estado)}.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao atualizar estado: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.distribuicao'))

@eleitoral_bp.route('/api/movimento/<int:id>/fluxo')
def api_movimento_fluxo(id):
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    param_marker = '%s' if is_pg else '?'
    try:
        c.execute(f"""
            SELECT m.*, lo.nome as origem, ld.nome as destino
            FROM eleitoral_movimento_material m
            JOIN eleitoral_local_armazenamento lo ON m.local_origem_id = lo.id
            JOIN eleitoral_local_armazenamento ld ON m.local_destino_id = ld.id
            WHERE m.id = {param_marker}
        """, (id,))
        mov = c.fetchone()
        if not mov:
            conn.close()
            return jsonify({'erro': 'Guia não encontrada'}), 404
        
        c.execute(f"""
            SELECT i.*, t.nome as material, t.variante
            FROM eleitoral_movimento_item i
            JOIN eleitoral_tipo_material t ON i.tipo_material_id = t.id
            WHERE i.movimento_id = {param_marker}
        """, (id,))
        itens = c.fetchall()
        
        c.execute(f"""
            SELECT h.*, u.nome_completo as utilizador_nome, u.username as utilizador_username
            FROM eleitoral_movimento_historico h
            LEFT JOIN users u ON h.utilizador_id = u.id
            WHERE h.movimento_id = {param_marker}
            ORDER BY h.id ASC
        """, (id,))
        historico = c.fetchall()
        
        estado_atual = mov['estado' if is_pg else 'estado']
        idx_atual = ESTADOS_DISTRIBUICAO_ORDEM.get(estado_atual, 0)
        anulada = estado_atual == 'ANULADA'
        
        user_perfil = session.get('perfil')
        user_local_id = session.get('eleitoral_local_id')
        is_admin = user_perfil == 'admin'
        origem_id = mov['local_origem_id' if is_pg else 'local_origem_id']
        destino_id = mov['local_destino_id' if is_pg else 'local_destino_id']
        pode_avancar = bool(is_admin or (user_local_id and origem_id == user_local_id)) and idx_atual < ESTADOS_DISTRIBUICAO_ORDEM['ENVIADO'] and not anulada
        pode_receber = idx_atual == ESTADOS_DISTRIBUICAO_ORDEM['ENVIADO'] and (is_admin or (user_local_id and destino_id == user_local_id)) and not anulada
        pode_anular = is_admin and not anulada
        proximos = [s for s, o in ESTADOS_DISTRIBUICAO_ORDEM.items() if o > idx_atual and o < ESTADOS_DISTRIBUICAO_ORDEM['RECEBIDO']] if not anulada else []
        
        conn.close()
        return jsonify({
            'movimento': dict(mov),
            'itens': [dict(i) for i in itens],
            'historico': [dict(h) for h in historico],
            'estados': ESTADOS_DISTRIBUICAO,
            'estado_atual': estado_atual,
            'pode_avancar': pode_avancar,
            'pode_receber': pode_receber,
            'pode_anular': pode_anular,
            'proximos': proximos,
        })
    except Exception as e:
        conn.close()
        return jsonify({'erro': str(e)}), 500

@eleitoral_bp.route('/locais')
def locais():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    c.execute("SELECT * FROM eleitoral_local_armazenamento ORDER BY CASE WHEN parent_id IS NULL THEN 0 ELSE 1 END, tipo, nome")
    locais_lista = c.fetchall()
    locais_row = [dict(l) for l in locais_lista] if not is_pg else locais_lista
    
    c.execute("SELECT id, nome FROM eleitoral_provincia ORDER BY nome")
    provincias = c.fetchall()
    
    # Mapa para mostrar o pai
    mapa_nomes = {l['id']: f"{l['tipo']} - {l['nome']}" for l in locais_lista}
    for l in locais_row:
        l['pai_nome'] = mapa_nomes.get(l['parent_id']) if l['parent_id'] else None
        l['n_filhos'] = sum(1 for x in locais_lista if x['parent_id'] == l['id'])
    
    conn.close()
    return render_template('eleitoral/locais.html', locais=locais_row, tipos=ELEITORAL_TIPOS_LOCAL, provincias=provincias)

@eleitoral_bp.route('/locais/criar', methods=['POST'])
def criar_local():
    nome = (request.form.get('nome') or '').strip()
    tipo = request.form.get('tipo')
    parent_id = request.form.get('parent_id') or None
    tem_filhos = 1 if request.form.get('tem_filhos') == 'sim' else 0
    provincia_id = request.form.get('provincia_id') or None
    observacoes = request.form.get('observacoes') or None
    
    if not nome or not tipo:
        flash("Nome e tipo do local são obrigatórios.", "error")
        return redirect(url_for('eleitoral.locais'))
    if tipo not in ELEITORAL_TIPOS_LOCAL:
        flash("Tipo de local inválido.", "error")
        return redirect(url_for('eleitoral.locais'))
    
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    try:
        param_marker = '%s' if is_pg else '?'
        c.execute(f"INSERT INTO eleitoral_local_armazenamento (nome, tipo, parent_id, tem_filhos, provincia_id, observacoes) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})",
                  (nome, tipo, parent_id, tem_filhos, provincia_id, observacoes))
        conn.commit()
        flash(f"Local '{nome}' criado com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao criar local: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.locais'))

@eleitoral_bp.route('/locais/editar/<int:id>', methods=['POST'])
def editar_local(id):
    nome = (request.form.get('nome') or '').strip()
    tipo = request.form.get('tipo')
    parent_id = request.form.get('parent_id') or None
    tem_filhos = 1 if request.form.get('tem_filhos') == 'sim' else 0
    provincia_id = request.form.get('provincia_id') or None
    observacoes = request.form.get('observacoes') or None
    activo = 1 if request.form.get('activo') == 'sim' else 0
    
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor()
    try:
        param_marker = '%s' if is_pg else '?'
        if parent_id and int(parent_id) == id:
            parent_id = None
        c.execute(f"UPDATE eleitoral_local_armazenamento SET nome = {param_marker}, tipo = {param_marker}, parent_id = {param_marker}, tem_filhos = {param_marker}, provincia_id = {param_marker}, observacoes = {param_marker}, activo = {param_marker} WHERE id = {param_marker}",
                  (nome, tipo, parent_id, tem_filhos, provincia_id, observacoes, activo, id))
        conn.commit()
        flash(f"Local '{nome}' atualizado com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao atualizar local: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('eleitoral.locais'))

@eleitoral_bp.route('/distribuicao')
def distribuicao():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    # Processo Ativo
    c.execute("SELECT * FROM eleitoral_processo_eleitoral WHERE estado = 'EM_CURSO' ORDER BY id DESC LIMIT 1")
    processo_ativo = c.fetchone()
    
    if not processo_ativo:
        conn.close()
        flash("Nenhum processo eleitoral ativo no momento. Inicie um processo para gerir a distribuição.", "warning")
        return redirect(url_for('eleitoral.dashboard'))
        
    processo_id = processo_ativo['id'] if is_pg else processo_ativo['id']
    
    # Locais
    c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 ORDER BY tipo, nome")
    locais = c.fetchall()
    
    # Materiais (Tipos)
    c.execute("SELECT * FROM eleitoral_tipo_material ORDER BY nome")
    tipos_material = c.fetchall()
    
    # Estoque Atual
    c.execute('''
        SELECT l.nome as local, t.nome as material, t.variante, s.quantidade_total as total, s.quantidade_bom as bom, s.quantidade_mau as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        WHERE s.processo_id = %s
        ORDER BY l.tipo, l.nome, t.nome
    ''' if is_pg else '''
        SELECT l.nome as local, t.nome as material, t.variante, s.quantidade_total as total, s.quantidade_bom as bom, s.quantidade_mau as mau
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        WHERE s.processo_id = ?
        ORDER BY l.tipo, l.nome, t.nome
    ''', (processo_id,))
    estoque = c.fetchall()
    
    # Guias de Movimentação (Em Trânsito e Recebidas)
    c.execute('''
        SELECT m.*, lo.nome as origem, ld.nome as destino
        FROM eleitoral_movimento_material m
        JOIN eleitoral_local_armazenamento lo ON m.local_origem_id = lo.id
        JOIN eleitoral_local_armazenamento ld ON m.local_destino_id = ld.id
        WHERE m.processo_id = %s
        ORDER BY m.id DESC LIMIT 50
    ''' if is_pg else '''
        SELECT m.*, lo.nome as origem, ld.nome as destino
        FROM eleitoral_movimento_material m
        JOIN eleitoral_local_armazenamento lo ON m.local_origem_id = lo.id
        JOIN eleitoral_local_armazenamento ld ON m.local_destino_id = ld.id
        WHERE m.processo_id = ?
        ORDER BY m.id DESC LIMIT 50
    ''', (processo_id,))
    movimentos_raw = c.fetchall()
    movimentos = [dict(m) for m in movimentos_raw] if not is_pg else movimentos_raw
    
    user_perfil = session.get('perfil')
    user_local_id = session.get('eleitoral_local_id')
    
    for m in movimentos:
        m_id = m['id']
        c.execute('''
            SELECT i.*, t.nome as material, t.variante
            FROM eleitoral_movimento_item i
            JOIN eleitoral_tipo_material t ON i.tipo_material_id = t.id
            WHERE i.movimento_id = %s
        ''' if is_pg else '''
            SELECT i.*, t.nome as material, t.variante
            FROM eleitoral_movimento_item i
            JOIN eleitoral_tipo_material t ON i.tipo_material_id = t.id
            WHERE i.movimento_id = ?
        ''', (m_id,))
        m['itens'] = c.fetchall()
        
        estado_atual = m['estado']
        idx_atual = ESTADOS_DISTRIBUICAO_ORDEM.get(estado_atual, 0)
        m['ordem_estado'] = idx_atual
        m['titulo_estado'] = TITULO_ESTADO.get(estado_atual, estado_atual)
        m['proximos'] = [s for s, o in ESTADOS_DISTRIBUICAO_ORDEM.items() if o > idx_atual and o < ESTADOS_DISTRIBUICAO_ORDEM['RECEBIDO']]
        
        origem_id = m['local_origem_id']
        destino_id = m['local_destino_id']
        is_admin = user_perfil == 'admin'
        is_origem = user_local_id is not None and origem_id == user_local_id
        is_destino = user_local_id is not None and (destino_id == user_local_id)
        m['anulada'] = estado_atual == 'ANULADA'
        m['pode_avancar'] = bool(is_admin or is_origem) and idx_atual < ESTADOS_DISTRIBUICAO_ORDEM['ENVIADO'] and not m['anulada']
        m['pode_receber'] = idx_atual == ESTADOS_DISTRIBUICAO_ORDEM['ENVIADO'] and (is_admin or is_destino) and not m['anulada']
        m['pode_anular'] = is_admin and not m['anulada']
    
    conn.close()
    
    return render_template('eleitoral/distribuicao.html', 
                           processo_ativo=processo_ativo,
                           locais=locais,
                           tipos_material=tipos_material,
                           estoque=estoque,
                           movimentos=movimentos,
                           titulos_estado=TITULO_ESTADO,
                           ordens_estado=ESTADOS_DISTRIBUICAO_ORDEM)

@eleitoral_bp.route('/distribuicao/aquisicao', methods=['POST'])
def nova_aquisicao():
    processo_id = request.form.get('processo_id')
    local_id = request.form.get('local_id')
    tipo_material_id = request.form.get('tipo_material_id')
    quantidade_bom = float(request.form.get('quantidade_bom') or 0)
    quantidade_mau = float(request.form.get('quantidade_mau') or 0)
    quantidade_total = quantidade_bom + quantidade_mau
    
    utilizador_id = session.get('user_id')
    
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    param_marker = '%s' if is_pg else '?'
    
    # Validação: a aquisição só é permitida num processo EM_CURSO do ano corrente ou futuro
    c.execute(f"SELECT * FROM eleitoral_processo_eleitoral WHERE id = {param_marker}", (processo_id,))
    processo = c.fetchone()
    if not processo:
        conn.close()
        flash("Processo eleitoral não encontrado.", "error")
        return redirect(url_for('eleitoral.distribuicao'))
    if processo['estado' if is_pg else 'estado'] != 'EM_CURSO':
        conn.close()
        flash("Aquisição recusada: o processo eleitoral não está ativo (EM CURSO).", "error")
        return redirect(url_for('eleitoral.distribuicao'))
    try:
        ano_processo = int(processo['ano' if is_pg else 'ano'])
    except (TypeError, ValueError):
        ano_processo = 0
    if ano_processo < datetime.datetime.now().year:
        conn.close()
        flash("Aquisição recusada: o processo pertence a um ano já decorrido.", "error")
        return redirect(url_for('eleitoral.distribuicao'))
    
    # Verifica se já existe registo deste material neste local para este processo
    if is_pg:
        c.execute("SELECT id FROM eleitoral_material_sobrante WHERE processo_id = %s AND local_id = %s AND tipo_material_id = %s", (processo_id, local_id, tipo_material_id))
    else:
        c.execute("SELECT id FROM eleitoral_material_sobrante WHERE processo_id = ? AND local_id = ? AND tipo_material_id = ?", (processo_id, local_id, tipo_material_id))
    
    row = c.fetchone()
    
    if row:
        row_id = row['id'] if is_pg else row['id']
        if is_pg:
            c.execute("UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total + %s, quantidade_bom = quantidade_bom + %s, quantidade_mau = quantidade_mau + %s WHERE id = %s",
                      (quantidade_total, quantidade_bom, quantidade_mau, row_id))
        else:
            c.execute("UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total + ?, quantidade_bom = quantidade_bom + ?, quantidade_mau = quantidade_mau + ? WHERE id = ?",
                      (quantidade_total, quantidade_bom, quantidade_mau, row_id))
    else:
        if is_pg:
            c.execute("INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id) VALUES (%s, %s, %s, %s, %s, %s, 'AQUISICAO', %s)",
                      (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, utilizador_id))
        else:
            c.execute("INSERT INTO eleitoral_material_sobrante (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id) VALUES (?, ?, ?, ?, ?, ?, 'AQUISICAO', ?)",
                      (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, utilizador_id))
            
    conn.commit()
    conn.close()
    flash("Nova aquisição registada com sucesso no inventário ativo!", "success")
    return redirect(url_for('eleitoral.distribuicao'))

@eleitoral_bp.route('/distribuicao/nova_guia', methods=['GET', 'POST'])
def nova_guia():
    conn, is_pg = get_eleitoral_db()
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    if request.method == 'POST':
        processo_id = request.form.get('processo_id')
        local_origem_id = request.form.get('local_origem_id')
        local_destino_id = request.form.get('local_destino_id')
        observacoes = request.form.get('observacoes')
        utilizador_id = session.get('user_id')
        
        # Get array of materials and quantities
        tipos_material_ids = request.form.getlist('item_tipo_material_id[]')
        quantidades_bom = request.form.getlist('item_quantidade_bom[]')
        quantidades_mau = request.form.getlist('item_quantidade_mau[]')
        
        try:
            # 1. Criar Guia (Movimento)
            param_marker = '%s' if is_pg else '?'
            if is_pg:
                c.execute("INSERT INTO eleitoral_movimento_material (processo_id, local_origem_id, local_destino_id, estado, observacoes_envio, utilizador_envio_id) VALUES (%s, %s, %s, 'EM_PREPARACAO', %s, %s) RETURNING id", 
                          (processo_id, local_origem_id, local_destino_id, observacoes, utilizador_id))
                mov_id = c.fetchone()['id']
            else:
                c.execute("INSERT INTO eleitoral_movimento_material (processo_id, local_origem_id, local_destino_id, estado, observacoes_envio, utilizador_envio_id) VALUES (?, ?, ?, 'EM_PREPARACAO', ?, ?)", 
                          (processo_id, local_origem_id, local_destino_id, observacoes, utilizador_id))
                mov_id = c.lastrowid

            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            # Regista o estado inicial no histórico
            c.execute(f"INSERT INTO eleitoral_movimento_historico (movimento_id, estado, observacoes, data, utilizador_id) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})",
                      (mov_id, 'EM_PREPARACAO', observacoes or 'Criação da guia', now, utilizador_id))
                
            # 2. Inserir itens e deduzir da origem
            for i in range(len(tipos_material_ids)):
                tm_id = tipos_material_ids[i]
                q_bom = float(quantidades_bom[i]) if quantidades_bom[i] else 0
                q_mau = float(quantidades_mau[i]) if quantidades_mau[i] else 0
                q_total = q_bom + q_mau
                
                if q_total > 0:
                    c.execute(f"INSERT INTO eleitoral_movimento_item (movimento_id, tipo_material_id, quantidade_bom, quantidade_mau) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker})", 
                              (mov_id, tm_id, q_bom, q_mau))
                    
                    # Deduz da origem
                    c.execute(f"UPDATE eleitoral_material_sobrante SET quantidade_total = quantidade_total - {param_marker}, quantidade_bom = quantidade_bom - {param_marker}, quantidade_mau = quantidade_mau - {param_marker} WHERE processo_id = {param_marker} AND local_id = {param_marker} AND tipo_material_id = {param_marker}",
                              (q_total, q_bom, q_mau, processo_id, local_origem_id, tm_id))
            
            conn.commit()
            flash("Guia de distribuição criada! O material está em fase de preparação.", "success")
        except Exception as e:
            conn.rollback()
            flash(f"Erro ao emitir guia: {str(e)}", "error")
            
        conn.close()
        return redirect(url_for('eleitoral.distribuicao'))
        
    # GET method
    c.execute("SELECT * FROM eleitoral_processo_eleitoral WHERE estado = 'EM_CURSO' ORDER BY id DESC LIMIT 1")
    processo_ativo = c.fetchone()
    
    if not processo_ativo:
        conn.close()
        flash("Nenhum processo ativo.", "warning")
        return redirect(url_for('eleitoral.dashboard'))
        
    user_perfil = session.get('perfil')
    user_local_id = session.get('eleitoral_local_id')
    
    if user_perfil == 'admin' or not user_local_id:
        c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 ORDER BY tipo, nome")
        locais_origem = c.fetchall()
        locais_destino = locais_origem
    else:
        if not is_pg:
            c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 AND id = ?", (user_local_id,))
        else:
            c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 AND id = %s", (user_local_id,))
        locais_origem = c.fetchall()
        
        if not is_pg:
            c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 AND (tipo = 'CENTRAL' OR parent_id = ? OR id = ?)", (user_local_id, user_local_id))
        else:
            c.execute("SELECT * FROM eleitoral_local_armazenamento WHERE activo = 1 AND (tipo = 'CENTRAL' OR parent_id = %s OR id = %s)", (user_local_id, user_local_id))
        locais_destino = c.fetchall()
        
    c.execute("SELECT * FROM eleitoral_tipo_material ORDER BY nome")
    tipos_material = c.fetchall()
    
    conn.close()
    return render_template('eleitoral/nova_guia.html', processo_ativo=processo_ativo, locais_origem=locais_origem, locais_destino=locais_destino, tipos_material=tipos_material)
    c = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if is_pg else conn.cursor()
    
    # Processo Ativo
    c.execute("SELECT id FROM eleitoral_processo_eleitoral WHERE estado = 'EM_CURSO' ORDER BY id DESC LIMIT 1")
    processo_ativo = c.fetchone()
    
    if not processo_ativo:
        conn.close()
        return jsonify({})
        
    processo_id = processo_ativo['id'] if is_pg else processo_ativo['id']
    
    # Map Provincias and Brigades directly associated
    param = (processo_id,)
    query = '''
        SELECT 
            p.nome as provincia, 
            t.nome as material,
            SUM(s.quantidade_total) as total
        FROM eleitoral_material_sobrante s
        JOIN eleitoral_local_armazenamento l ON s.local_id = l.id
        JOIN eleitoral_provincia p ON l.provincia_id = p.id OR (l.parent_id IS NOT NULL AND l.provincia_id IS NULL)
        JOIN eleitoral_tipo_material t ON s.tipo_material_id = t.id
        WHERE s.processo_id = ?
        GROUP BY p.nome, t.nome
    '''
    if is_pg:
        query = query.replace('?', '%s')
        
    try:
        c.execute(query, param)
        rows = c.fetchall()
        
        # Build dictionary { 'Maputo': { 'Impressoras': 100, 'Kits': 50 }, 'Gaza': ... }
        res = {}
        for r in rows:
            prov = r['provincia'] if is_pg else r['provincia']
            mat = r['material'] if is_pg else r['material']
            tot = r['total'] if is_pg else r['total']
            
            # normalize province names for map keys (e.g. Maputo, Cabo Delgado)
            prov_key = prov.strip()
            
            if prov_key not in res:
                res[prov_key] = {}
            res[prov_key][mat] = tot
            
        conn.close()
        return jsonify(res)
    except Exception as e:
        conn.close()
        return jsonify({"error": str(e)})

