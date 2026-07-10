import sqlite3
import docx

def insert_catalogs_and_process():
    from app import DB_PATH
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # 1. Insert Categories if missing
    cats = [
        (1, 'Equipamento Electrónico'),
        (2, 'Equipamento de Identificação/Impressão'),
        (3, 'Consumível'),
        (4, 'Mobiliário e Acessório')
    ]
    for cat in cats:
        c.execute("INSERT OR IGNORE INTO eleitoral_categoria_material (id, nome) VALUES (?, ?)", cat)
        
    # 2. Insert Tipos de Material
    tipos = [
        (1, 1, 'Mobile ID 23/24', None, 'Unidade', 1),
        (2, 1, 'Mobile ID 18/19', None, 'Unidade', 1),
        (3, 2, 'Impressora PVC', None, 'Unidade', 1),
        (4, 4, 'Tripé/Pano de fundo', None, 'Unidade', 1),
        (5, 2, 'Cartões PVC', None, 'Unidade', 0),
        (6, 3, 'Boletim de Inscrição', None, 'Unidade', 0),
        (7, 1, 'Painel Solar', None, 'Unidade', 1),
        (8, 4, 'Mochila', None, 'Unidade', 1),
        (9, 4, 'Mochila', 'Com Painel', 'Unidade', 1),
        (10, 4, 'Mochila', 'Sem Painel', 'Unidade', 0),
        (11, 4, 'Malas Metálicas', None, 'Unidade', 1),
        (12, 3, 'Toner de Impressão de Cadernos', None, 'Unidade', 0),
        (13, 3, 'Toner para Impressão do Cartão', None, 'Unidade', 0),
        (14, 3, 'Kit de Limpeza', None, 'Unidade', 0)
    ]
    for tipo in tipos:
        c.execute("""
            INSERT OR IGNORE INTO eleitoral_tipo_material 
            (id, categoria_id, nome, variante, unidade_medida, controla_estado) 
            VALUES (?, ?, ?, ?, ?, ?)
        """, tipo)
        
    # 3. Create initial process
    c.execute("SELECT id FROM eleitoral_processo_eleitoral WHERE nome = 'Recenseamento Eleitoral 2024'")
    proc = c.fetchone()
    if not proc:
        c.execute("INSERT INTO eleitoral_processo_eleitoral (nome, tipo, ano, estado) VALUES (?, ?, ?, ?)",
                 ('Recenseamento Eleitoral 2024', 'Recenseamento', 2024, 'EM_CURSO'))
        proc_id = c.lastrowid
    else:
        proc_id = proc[0]
        
    conn.commit()
    conn.close()
    return proc_id

def parse_num(val):
    if not val or val == '-': return 0
    try:
        val = val.replace(',', '.')
        return float(val)
    except:
        return 0

def import_docx(proc_id, docx_path):
    from app import DB_PATH
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    doc = docx.Document(docx_path)
    
    # Cache local ids
    c.execute("SELECT id, nome, tipo FROM eleitoral_local_armazenamento")
    locais = {row[1].strip().lower(): row[0] for row in c.fetchall()}
    
    # Custom mapping for local names from doc
    local_map = {
        'c. de maputo': locais.get('cidade de maputo', 1),
        'maputo': locais.get('maputo', 2),
        'gaza': locais.get('gaza', 3),
        'inhambane': locais.get('inhambane', 4),
        'sofala': locais.get('sofala', 5),
        'manica': locais.get('manica', 6),
        'tete': locais.get('tete', 7),
        'zambézia': locais.get('zambézia', 8),
        'nampula': locais.get('nampula', 9),
        'cabo delgado': locais.get('cabo delgado', 10),
        'niassa': locais.get('niassa', 11),
        'áfrica do sul': locais.get('áfrica do sul'),
        'eswatini': locais.get('eswatini'),
        'zimbabwe': locais.get('zimbabwe'),
        'malawi': locais.get('malawi'),
        'zâmbia': locais.get('zâmbia'),
        'tanzânia': locais.get('tanzânia'),
        'quénia': locais.get('quénia'),
        'portugal': locais.get('portugal'),
        'alemanha': locais.get('alemanha')
    }
    
    # Table 1: Provinces
    t1 = doc.tables[0]
    for j, row in enumerate(t1.rows):
        if j < 2: continue # skip header
        cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
        if not cells[1]: continue
        
        local_name = cells[1].lower()
        if 'total' in local_name: continue
        
        local_id = local_map.get(local_name)
        if not local_id: continue
        
        # Insert records for Table 1
        inserts = [
            (1, parse_num(cells[2]), parse_num(cells[3]), parse_num(cells[4])), # Mobile ID 23/24
            (2, parse_num(cells[5]), parse_num(cells[6]), parse_num(cells[7])), # Mobile ID 18/19
            (3, parse_num(cells[8]), parse_num(cells[9]), parse_num(cells[10])), # Impressora PVC
            (4, parse_num(cells[11]), 0, 0), # Tripé
            (5, parse_num(cells[12]), 0, 0), # Cartões PVC
            (6, parse_num(cells[13]), 0, 0), # Boletim
            (7, parse_num(cells[14]), 0, 0), # Painel Solar
            (8, parse_num(cells[15]), 0, 0), # Mochila (sem variante)
            (11, parse_num(cells[16]), parse_num(cells[17]), parse_num(cells[18])), # Malas Metálicas
            (12, parse_num(cells[19]), 0, 0), # Toner Cadernos
            (13, parse_num(cells[20]), 0, 0), # Toner Cartão
            (14, parse_num(cells[21]), 0, 0), # Kit Limpeza
        ]
        for tipo_id, total, bom, mau in inserts:
            if total > 0 or bom > 0 or mau > 0:
                # Se o total for 0 mas tem bom ou mau, assumimos que total = bom + mau
                if total == 0: total = bom + mau
                c.execute("""
                    INSERT OR REPLACE INTO eleitoral_material_sobrante 
                    (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id)
                    VALUES (?, ?, ?, ?, ?, ?, 'IMPORTACAO_EXCEL', 1)
                """, (proc_id, local_id, tipo_id, total, bom, mau))

    # Table 2: Diaspora
    t2 = doc.tables[1]
    for j, row in enumerate(t2.rows):
        if j < 3: continue
        cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
        if not cells[0]: continue
        
        local_name = cells[0].lower()
        if 'total' in local_name: continue
        
        local_id = local_map.get(local_name)
        if not local_id: continue
        
        # Mappings for diaspora
        # 1: Mobile ID (Quant, Mau, Bom)
        # 4: Impressoras PVC (Quant, Mau, Bom) -> Wait, cells 4, 5, 6
        # 7: Tripé (Quant, Mau, Bom) -> cells 7, 8, 9
        # 10: Mochilas com painel (Quant, Mau, Bom) -> cells 10, 11, 12
        # 13: Mochilas sem painel (Quant) -> cell 13
        
        m_total = parse_num(cells[1])
        m_mau = parse_num(cells[2])
        m_bom = parse_num(cells[3])
        if m_total == 0: m_total = m_bom + m_mau
        
        i_total = parse_num(cells[4])
        i_mau = parse_num(cells[5])
        i_bom = parse_num(cells[6])
        if i_total == 0: i_total = i_bom + i_mau
        
        t_total = parse_num(cells[7])
        t_mau = parse_num(cells[8])
        t_bom = parse_num(cells[9])
        if t_total == 0: t_total = t_bom + t_mau
        
        mc_total = parse_num(cells[10])
        mc_mau = parse_num(cells[11])
        mc_bom = parse_num(cells[12])
        if mc_total == 0: mc_total = mc_bom + mc_mau
        
        ms_total = parse_num(cells[13])
        
        inserts = [
            (1, m_total, m_bom, m_mau),
            (3, i_total, i_bom, i_mau),
            (4, t_total, t_bom, t_mau),
            (9, mc_total, mc_bom, mc_mau),
            (10, ms_total, 0, 0)
        ]
        for tipo_id, total, bom, mau in inserts:
            if total > 0 or bom > 0 or mau > 0:
                c.execute("""
                    INSERT OR REPLACE INTO eleitoral_material_sobrante 
                    (processo_id, local_id, tipo_material_id, quantidade_total, quantidade_bom, quantidade_mau, origem, utilizador_id)
                    VALUES (?, ?, ?, ?, ?, ?, 'IMPORTACAO_EXCEL', 1)
                """, (proc_id, local_id, tipo_id, total, bom, mau))

    conn.commit()
    conn.close()
    print("Importacao concluida!")

if __name__ == '__main__':
    pid = insert_catalogs_and_process()
    import_docx(pid, r"c:\Users\Acer\Documents\tecnologias\ddgeiinout\Material_Recenseamento_Eleitoral_A3 (2).docx")
